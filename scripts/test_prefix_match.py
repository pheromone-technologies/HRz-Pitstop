import openpyxl
import json
import re

wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))

final_map_by_name = {}
final_map_by_sku = {}

for row in sheet.iter_rows(values_only=True):
    d = dict(zip(headers, row))
    pname = (d.get('Product Name') or '').strip().lower()
    sku = (d.get('SKU') or '').strip()
    cat = (d.get('Category') or '').strip()
    subcat = (d.get('Subcategory') or '').strip()
    if pname:
        final_map_by_name[pname] = (cat, subcat)
    if sku and sku != 'Missing SKU':
        final_map_by_sku[sku] = (cat, subcat)
wb.close()

with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

matched = 0
unmatched_list = []

for p in products:
    pname = p.get('name', '').strip().lower()
    sku = p.get('sku', '').strip()
    
    cat_sub = None
    if pname in final_map_by_name:
        cat_sub = final_map_by_name[pname]
    elif sku in final_map_by_sku:
        cat_sub = final_map_by_sku[sku]
    else:
        # Check if any variant matches
        if 'variants' in p:
            for v in p['variants']:
                vname = (v.get('name') or '').strip().lower()
                vsku = (v.get('sku') or '').strip()
                if vname in final_map_by_name:
                    cat_sub = final_map_by_name[vname]
                    break
                elif vsku in final_map_by_sku:
                    cat_sub = final_map_by_sku[vsku]
                    break
    
    if not cat_sub:
        # Check prefix match in final_map_by_name
        for fn_name, val in final_map_by_name.items():
            if fn_name.startswith(pname) or pname.startswith(fn_name):
                cat_sub = val
                break
                
    if cat_sub:
        matched += 1
    else:
        unmatched_list.append(p)

print(f"Total storefront products: {len(products)}")
print(f"Matched with variant / prefix support: {matched}")
print(f"Remaining unmatched: {len(unmatched_list)}")
for p in unmatched_list[:5]:
    print(f"  Unmatched: {p['name']} (Cat: {p.get('category')})")
