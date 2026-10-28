import openpyxl
import json
import os
import shutil

src_final = 'Final_Master_Catalogue_VERIFIED.xlsx'
bak_final = 'Final_Master_Catalogue_VERIFIED.xlsx.bak'

# 1. Create a safe backup before any modification
shutil.copy2(src_final, bak_final)
print(f"Backed up {src_final} -> {bak_final}")

# 2. Build image mapping
wb_old = openpyxl.load_workbook('Verified_Products_FINAL-1.xlsx', data_only=True)
sheet_old = wb_old['Verified Products']
headers_old = [str(c).strip() for c in next(sheet_old.iter_rows(min_row=1, max_row=1, values_only=True))]

img_idx = headers_old.index('Product Image Filename / URL')
gal_idx = headers_old.index('Gallery Images')
name_idx = headers_old.index('Product Name')
sku_idx = headers_old.index('SKU')

old_img_map_by_name = {}
old_img_map_by_sku = {}

for row in sheet_old.iter_rows(min_row=2, values_only=True):
    pname = str(row[name_idx]).strip() if row[name_idx] else ''
    sku = str(row[sku_idx]).strip() if row[sku_idx] else ''
    img = row[img_idx] or ''
    gal = row[gal_idx] or ''
    if pname and (img or gal):
        old_img_map_by_name[pname.lower()] = (img, gal)
    if sku and sku != 'Missing SKU' and (img or gal):
        old_img_map_by_sku[sku] = (img, gal)
wb_old.close()

# Also load from products.json
with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

json_img_map = {}
for p in products:
    pname = p.get('name', '').strip().lower()
    img = p.get('image', '')
    gal = ', '.join(p.get('gallery', [])) if isinstance(p.get('gallery'), list) else (p.get('gallery') or '')
    if pname and (img or gal):
        json_img_map[pname] = (img, gal)
    if 'variants' in p:
        for v in p['variants']:
            vname = (v.get('name') or '').strip().lower()
            vsku = (v.get('sku') or '').strip()
            vimg = v.get('image', img)
            vgal = ', '.join(v.get('gallery', [])) if isinstance(v.get('gallery'), list) else gal
            if vname:
                json_img_map[vname] = (vimg, vgal)
            if vsku and vsku != 'Missing SKU':
                old_img_map_by_sku[vsku] = (vimg, vgal)

# 3. Open Final_Master_Catalogue_VERIFIED.xlsx and append columns
wb_final = openpyxl.load_workbook(src_final)
sheet_final = wb_final['Master Catalogue']

headers_final = [str(cell.value).strip() if cell.value is not None else '' for cell in sheet_final[1]]
print("Current headers in Master Catalogue:", len(headers_final))

if 'Product Image Filename / URL' not in headers_final:
    col_img = len(headers_final) + 1
    col_gal = len(headers_final) + 2
    sheet_final.cell(row=1, column=col_img, value='Product Image Filename / URL')
    sheet_final.cell(row=1, column=col_gal, value='Gallery Images')
else:
    col_img = headers_final.index('Product Image Filename / URL') + 1
    col_gal = headers_final.index('Gallery Images') + 1

fn_name_col = headers_final.index('Product Name') + 1
fn_sku_col = headers_final.index('SKU') + 1

matched_img_count = 0
matched_gal_count = 0

for row_idx in range(2, sheet_final.max_row + 1):
    pname_cell = sheet_final.cell(row=row_idx, column=fn_name_col).value
    sku_cell = sheet_final.cell(row=row_idx, column=fn_sku_col).value
    
    pname = str(pname_cell).strip().lower() if pname_cell else ''
    sku = str(sku_cell).strip() if sku_cell else ''
    
    img_gal = None
    if pname in old_img_map_by_name:
        img_gal = old_img_map_by_name[pname]
    elif sku in old_img_map_by_sku:
        img_gal = old_img_map_by_sku[sku]
    elif pname in json_img_map:
        img_gal = json_img_map[pname]
    else:
        # Prefix match
        for k, v in json_img_map.items():
            if pname.startswith(k) or k.startswith(pname):
                img_gal = v
                break

    if img_gal:
        img, gal = img_gal
        if img:
            sheet_final.cell(row=row_idx, column=col_img, value=img)
            matched_img_count += 1
        if gal:
            sheet_final.cell(row=row_idx, column=col_gal, value=gal)
            matched_gal_count += 1

wb_final.save(src_final)
wb_final.close()
print(f"Successfully enriched Final_Master_Catalogue_VERIFIED.xlsx!")
print(f"  Images attached: {matched_img_count}")
print(f"  Galleries attached: {matched_gal_count}")
