import openpyxl
import json

wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))

final_map_by_name = {}
final_map_by_sku = {}

for row in sheet.iter_rows(values_only=True):
    d = dict(zip(headers, row))
    pname = (d.get('Product Name') or '').strip().lower()
    sku = (d.get('SKU') or '').strip()
    cat = (d.get('Category') or '').strip()
    subcat = (d.get('Subcategory') or '').strip()
    if pname:
        final_map_by_name[pname] = (cat, subcat)
    if sku and sku != 'Missing SKU':
        final_map_by_sku[sku] = (cat, subcat)
wb.close()

with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

updated_cat_count = 0
updated_subcat_count = 0
no_match_count = 0

for p in products:
    pname = p.get('name', '').strip().lower()
    sku = p.get('sku', '').strip()
    
    new_cat_sub = None
    if pname in final_map_by_name:
        new_cat_sub = final_map_by_name[pname]
    elif sku in final_map_by_sku:
        new_cat_sub = final_map_by_sku[sku]
    
    if new_cat_sub:
        old_cat = p.get('category', '')
        old_subcat = p.get('subcategory', '')
        new_cat, new_subcat = new_cat_sub
        if new_cat and old_cat != new_cat:
            updated_cat_count += 1
        if new_subcat and old_subcat != new_subcat:
            updated_subcat_count += 1
    else:
        no_match_count += 1

print(f"Total storefront products: {len(products)}")
print(f"Products with Category updated: {updated_cat_count}")
print(f"Products with Subcategory updated: {updated_subcat_count}")
print(f"Products without match in Master: {no_match_count}")
