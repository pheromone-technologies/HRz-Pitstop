import openpyxl
import os
import sys

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"E:\coad\hrx website 2"
FINAL_XLSX = os.path.join(BASE_DIR, "HRZ_Pitstop_Client_Inventory_FINAL.xlsx")
GOLD_XLSX = os.path.join(BASE_DIR, "HRZ_Gold_Master_Verified.xlsx")

wb_final = openpyxl.load_workbook(FINAL_XLSX, data_only=True)
wb_gold = openpyxl.load_workbook(GOLD_XLSX, data_only=True)

# Build Gold image lookup by normalized name
gold_images = {}
for s in wb_gold.sheetnames:
    if s in ['LISTS', 'BIKE LISTS', 'MODEL DROPDOWNS']:
        continue
    ws = wb_gold[s]
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        nm = str(r[0]).strip().lower()
        img = r[9] if len(r) > 9 else None
        if img and str(img).strip() and str(img).strip() != 'None':
            gold_images[nm] = str(img).strip()

print(f"Loaded {len(gold_images)} verified images from Gold Master.")

product_sheets = [s for s in wb_final.sheetnames if s not in ['LISTS', 'BIKE LISTS', 'MODEL DROPDOWNS']]

total_final = 0
final_has_own_img = 0
final_own_img_exists = 0
final_can_inherit_gold = 0
final_missing_img = 0

for s in product_sheets:
    ws = wb_final[s]
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not any(c is not None and str(c).strip() != '' for c in r):
            continue
        total_final += 1
        name = str(r[0] or '').strip()
        img = str(r[9] or '').strip() if len(r) > 9 else ''
        
        if img and img != 'None':
            final_has_own_img += 1
            full_img_path = os.path.join(BASE_DIR, img.replace('/', os.sep))
            if os.path.exists(full_img_path) and os.path.getsize(full_img_path) > 500:
                final_own_img_exists += 1
            else:
                # print(f"Missing disk image: {img}")
                pass
        else:
            # Check if in Gold
            nm_lower = name.lower()
            if nm_lower in gold_images:
                final_can_inherit_gold += 1
            else:
                final_missing_img += 1

print(f"Total rows in FINAL:                {total_final}")
print(f"Has own image specified:            {final_has_own_img}")
print(f"Own image verified existing on disk:{final_own_img_exists}")
print(f"Can inherit from Gold Master:       {final_can_inherit_gold}")
print(f"Missing image entirely:             {final_missing_img}")
print(f"Total with verified images available: {final_own_img_exists + final_can_inherit_gold}")
