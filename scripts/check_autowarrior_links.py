import openpyxl

wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))

count = 0
for row in sheet.iter_rows(values_only=True):
    d = dict(zip(headers, row))
    ev = str(d.get('Evidence Link') or '')
    if 'autowarrior' in ev.lower():
        count += 1
        print(f"Product: {d.get('Product Name')} | Link: {ev} | Img: {d.get('Product Image Filename / URL')}")
wb.close()
print(f"Total autowarrior links: {count}")
