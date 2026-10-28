import openpyxl

wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
for sname in wb.sheetnames:
    sheet = wb[sname]
    for idx, row in enumerate(sheet.iter_rows(values_only=True)):
        row_str = ' '.join(str(c) for c in row if c is not None)
        if 'warrior' in row_str.lower():
            print(f"Sheet [{sname}] Row {idx+1}: {row_str[:120]}")
wb.close()
