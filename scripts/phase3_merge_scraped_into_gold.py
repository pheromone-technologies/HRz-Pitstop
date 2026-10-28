#!/usr/bin/env python3
"""
Phase 3: Merge Verified Scraped Data into Gold Master & Reconcile
Repo: HRz Pitstop (E:\\coad\\hrx website 2)

Merges:
- 100% verified items from scraped_verified_results.json into HRZ_Gold_Master_Verified.xlsx
- Validates image file exists on disk and > 1KB
- Validates selling price > purchase price (when purchase price > 0)
- Moves any non-conforming items to scrape_manual_review.csv
- Reconciles 100% of the 2,502 POS inventory items:
    Gold Master Verified (Original 122 + Newly Scraped 1005 = 1127)
  + Client Hold (406)
  + Scrape Manual Review (969)
  = 2,502 / 2,502 (100.0%)
"""

import os
import sys
import csv
import json
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# Force UTF-8 stdout
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
GOLD_XLSX = os.path.join(BASE_DIR, 'HRZ_Gold_Master_Verified.xlsx')
RESULTS_JSON = os.path.join(BASE_DIR, 'scraped_verified_results.json')
if not os.path.exists(RESULTS_JSON):
    RESULTS_JSON = os.path.join(BASE_DIR, '_vault_archive', 'scraper_working_cache', 'scraped_verified_results.json')

QUEUE_CSV = os.path.join(BASE_DIR, 'Needs_Scraping_Queue.csv')
if not os.path.exists(QUEUE_CSV):
    QUEUE_CSV = os.path.join(BASE_DIR, '_vault_archive', 'scraper_working_cache', 'Needs_Scraping_Queue.csv')

MANUAL_CSV = os.path.join(BASE_DIR, 'scrape_manual_review.csv')
if not os.path.exists(MANUAL_CSV):
    MANUAL_CSV = os.path.join(BASE_DIR, '_vault_archive', 'scraper_working_cache', 'scrape_manual_review.csv')

HOLD_XLSX = os.path.join(BASE_DIR, 'Client_Hold_Generic_And_Unfound.xlsx')
if not os.path.exists(HOLD_XLSX):
    HOLD_XLSX = os.path.join(BASE_DIR, '_vault_archive', 'legacy_catalogs_and_workbooks', 'Client_Hold_Generic_And_Unfound.xlsx')

MASTER_XLSX = os.path.join(BASE_DIR, 'Verified_Stock_Inventory_Master_Current.xlsx')
if not os.path.exists(MASTER_XLSX):
    MASTER_XLSX = os.path.join(BASE_DIR, 'products data', 'Verified_Stock_Inventory_Master_Current.xlsx')

VALID_SCHEMA = {
    'Helmets & Accessories': [
        'Full-Face Helmet', 'Open-Face Helmet', 'Modular Helmet', 'Offroad Helmet',
        'Helmet Visor', 'Helmet Accessories', 'Communication & Intercoms', 'Riding Eyewear & Goggles'
    ],
    'Riding Gears': [
        'Riding Jackets', 'Riding Pants', 'Riding Suits', 'Riding Gloves',
        'Riding Boots & Shoes', 'Rain Wear', 'Body Armor & Protectors', 'Thermal & Base Layers'
    ],
    'Bike Protection': [
        'Crash Guards & Sliders', 'Engine & Sump Guards', 'Radiator Guards',
        'Headlight Grills & Protectors', 'Fork & Axle Sliders', 'Handguards & Lever Protectors',
        'Brake Reservoir Guards', 'Pillion Footrest Protection'
    ],
    'Handlebar & Accessories': [
        'Handlebars & Clip-ons', 'Handlebar Risers', 'Grips and Throttle',
        'Bar End Weights', 'Mirrors & Bar-End Mirrors', 'Brake & Clutch Levers',
        'Master Cylinder Covers', 'Lever Guards'
    ],
    'Foot Rests': [
        'Rider Footpegs', 'Pillion Footrests', 'Footrest Hangers & Brackets',
        'Highway Pegs', 'Rearsets & Linkages'
    ],
    'Oil Caps & Stands': [
        'Engine Oil Filler Caps', 'Paddock Stands & Bobbins', 'Main Stands & Centre Stands',
        'Side Stand Extenders', 'Oil Filter Covers'
    ],
    'Body Fairing and Fenders': [
        'Front Fenders & Mudguards', 'Rear Huggers & Tyre Huggers', 'Visors & Windshields',
        'Belly Pans & Engine Cowls', 'Tail Tidies & Number Plate Holders',
        'Seat Cowls & Fairing Panels', 'Tank Covers & Extensions'
    ],
    'Bike Essentials': [
        'Mobile Holders & Mounts', 'GPS & Navigation Mounts', 'USB Chargers & Sockets',
        'Tyre Pressure Monitors (TPMS)', 'Bungee Cords & Cargo Nets', 'Bike Covers & Dust Sheets'
    ],
    'Cleaning Accessories': [
        'Chain Cleaners & Degreasers', 'Chain Lubes & Sprays', 'Shampoos & Washes',
        'Polish, Wax & Detailers', 'Microfiber Cloths & Sponges', 'Cleaning Brushes & Tools'
    ],
    'Performance Accessories': [
        'Performance Exhausts & Slip-ons', 'Air Filters & Intakes', 'Braking Components',
        'Suspension & Lowering Kits', 'Chain and Sprocket', 'ECU & Performance Tuning',
        'Spark Plugs & Ignition'
    ],
    'Lighting': [
        'LED Headlight Bulbs', 'Auxiliary & Fog Lights', 'Turn Signals & Indicators',
        'Tail Lights & Brake Lights', 'Light Mounts & Clamps', 'Flashers & Wiring Harnesses'
    ],
    'Guards & Mounts': [
        'Exhaust Heat Shields', 'Silencer Guards', 'Caliper Guards',
        'Sensor & Wire Protectors', 'GoPro & Action Cam Mounts', 'Number Plate Brackets'
    ],
    'Luggage & Touring': [
        'Top Boxes', 'Saddlebags & Panniers', 'Tank Bags', 'Tail Bags & Duffle Bags',
        'Backpacks & Hydration Packs', 'Top Racks & Luggage Carriers',
        'Saddle Stays & Pannier Mounts', 'Jerry Cans & Fuel Carriers'
    ],
    'General Accessories': [
        'Keychains & Key Covers', 'Stickers, Decals & Badges', 'Tank Pads & Grip Pads',
        'Merchandise & Decals', 'Locks & Anti-Theft', 'Tools & Workshop', 'Miscellaneous'
    ]
}

def normalize_subcategory(cat, subcat, product_name=''):
    valid_subs = VALID_SCHEMA.get(cat, [])
    if subcat in valid_subs:
        return subcat

    # Fuzzy match within category
    s_lower = str(subcat or '').lower().strip()
    for vs in valid_subs:
        if vs.lower() == s_lower or vs.lower() in s_lower:
            return vs

    # Category defaults
    if cat == 'General Accessories':
        return 'Miscellaneous'
    if cat == 'Cleaning Accessories':
        return 'Cleaning Brushes & Tools'
    if cat == 'Bike Essentials':
        return 'Mobile Holders & Mounts'

    return valid_subs[-1] if valid_subs else 'Miscellaneous'

def count_actual_products(ws):
    """Count rows in a worksheet that have non-empty Column A (excluding header)."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) <= 1:
        return 0
    return sum(1 for r in rows[1:] if r and r[0] is not None and str(r[0]).strip() != '')

def main():
    print("=" * 80)
    print("PHASE 3: MERGE VERIFIED SCRAPED DATA INTO GOLD MASTER & RECONCILE")
    print("=" * 80)

    # 1. Load Scraped Results
    with open(RESULTS_JSON, 'r', encoding='utf-8') as f:
        scraped_items = json.load(f)
    print(f"Loaded {len(scraped_items)} items from {RESULTS_JSON}.")

    # 2. Strict Image & Price Validation Gate
    verified_to_merge = []
    failed_scraped = []

    for it in scraped_items:
        img_p = it.get('image_path')
        has_real_img = False
        if img_p:
            full_img_p = os.path.join(BASE_DIR, img_p) if not os.path.isabs(img_p) else img_p
            if os.path.exists(full_img_p) and os.path.getsize(full_img_p) > 1024:
                has_real_img = True

        sp = it.get('scraped_selling_price', 0)
        pp = it.get('pos_purchase_price', 0)
        price_ok = (sp > pp) if pp > 0 else (sp > 0)

        if has_real_img and price_ok:
            verified_to_merge.append(it)
        else:
            reason = []
            if not has_real_img:
                reason.append("IMAGE_FILE_MISSING_OR_EMPTY")
            if not price_ok:
                reason.append(f"PRICE_INVALID (Scraped ₹{sp} <= Purchase ₹{pp})")
            failed_scraped.append((it, ", ".join(reason)))

    print(f"-> 100% Validated items ready to merge into Gold Master: {len(verified_to_merge)}")
    print(f"-> Scraped items failing image/price gate (sent to manual review): {len(failed_scraped)}")

    # 3. Load Existing Gold Master
    wb_gold = openpyxl.load_workbook(GOLD_XLSX)
    existing_gold_counts = {}
    for s_name in wb_gold.sheetnames[:14]:
        ws = wb_gold[s_name]
        existing_gold_counts[s_name] = count_actual_products(ws)
    total_existing_gold = sum(existing_gold_counts.values())
    print(f"\nExisting items in HRZ_Gold_Master_Verified.xlsx: {total_existing_gold}")

    # Formatting definitions
    data_font = Font(name='Segoe UI', size=10)
    border_thin = Side(border_style='thin', color='FFE0E0E0')
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    # 4. Check if already merged or append
    merged_by_cat = {s: 0 for s in wb_gold.sheetnames[:14]}
    merged_source_rows = set()

    # Collect existing product names or SKUs to avoid double-appending if re-run
    existing_names_by_sheet = {}
    for s_name in wb_gold.sheetnames[:14]:
        ws = wb_gold[s_name]
        existing_names_by_sheet[s_name] = {str(r[0]).strip().lower() for r in ws.iter_rows(values_only=True) if r and r[0]}

    for it in verified_to_merge:
        source_row = it['source_row']
        merged_source_rows.add(source_row)

        target_cat = it['category']
        if target_cat not in wb_gold.sheetnames:
            target_cat = 'General Accessories'

        ws = wb_gold[target_cat]
        pname = it.get('official_title') or it.get('pos_product_name')

        # If already exists in sheet, record it and continue
        if str(pname).strip().lower() in existing_names_by_sheet[target_cat]:
            merged_by_cat[target_cat] += 1
            continue

        subcat = normalize_subcategory(target_cat, it.get('subcategory', ''), it.get('pos_product_name', ''))

        mrp = it.get('scraped_mrp', 0)
        sp = it.get('scraped_selling_price', 0)
        discount = round(((mrp - sp) / mrp) * 100.0, 1) if mrp > sp and mrp > 0 else 0.0

        brand_val = it.get('brand') or it.get('vendor') or 'HRz Pitstop'
        stock_val = it.get('stock_qty') or 1
        img_val = it.get('image_path') or ''
        desc_val = it.get('description') or f"{pname} - Premium motorcycle gear and accessories verified by {brand_val}."
        sku_val = it.get('scraped_sku') or f"HRZ-{it['source_row'].replace(' ', '')}"

        # Find the next empty row specifically in Column A
        next_r = 2
        while ws.cell(row=next_r, column=1).value is not None:
            next_r += 1

        row_data = [
            pname, subcat, brand_val, 'Universal', 'Universal',
            int(mrp), int(sp), discount, stock_val,
            img_val, desc_val, sku_val, '', ''
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=next_r, column=col_idx, value=val)
            cell.font = data_font
            cell.border = cell_border
            if col_idx in [6, 7]:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx in [8, 9]:
                cell.alignment = Alignment(horizontal='right')
            elif col_idx in [2, 3, 4, 5, 12, 13, 14]:
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.alignment = Alignment(horizontal='left')

        existing_names_by_sheet[target_cat].add(str(pname).strip().lower())
        merged_by_cat[target_cat] += 1

    # 5. Ensure Data Validations are fully applied to all sheets
    for cat_name, valid_subs in VALID_SCHEMA.items():
        if cat_name in wb_gold.sheetnames:
            ws = wb_gold[cat_name]
            subs_str = '"' + ','.join(valid_subs) + '"'
            dv_sub = DataValidation(type='list', formula1=subs_str, allow_blank=True)
            dv_make = DataValidation(type='list', formula1='=$P$2:$P$100', allow_blank=True)
            dv_model = DataValidation(type='list', formula1='=$Q$2:$Q$200', allow_blank=True)

            ws.add_data_validation(dv_sub)
            ws.add_data_validation(dv_make)
            ws.add_data_validation(dv_model)

            dv_sub.add('B2:B2000')
            dv_make.add('D2:D2000')
            dv_model.add('E2:E2000')

    wb_gold.save(GOLD_XLSX)
    print(f"\nSaved updated {GOLD_XLSX} with {len(verified_to_merge)} newly scraped items!")

    # 6. Reconcile and write updated scrape_manual_review.csv
    with open(QUEUE_CSV, 'r', encoding='utf-8') as f:
        queue_rows = list(csv.DictReader(f))

    manual_reasons = {}
    if os.path.exists(MANUAL_CSV):
        with open(MANUAL_CSV, 'r', encoding='utf-8') as f:
            for r in csv.DictReader(f):
                manual_reasons[r['Source_Row']] = r.get('Reason', 'NO_MATCH_OR_DEAD_URL')

    for it, r in failed_scraped:
        manual_reasons[it['source_row']] = r

    final_manual_review_rows = []
    for q in queue_rows:
        s_row = q['Source_Row']
        if s_row not in merged_source_rows:
            reason = manual_reasons.get(s_row, 'NO_MATCH_OR_DEAD_URL')
            final_manual_review_rows.append({
                'Source_Row': s_row,
                'POS_Product_Name': q['POS_Product_Name'],
                'Brand': q['Brand'],
                'Category': q['Category'],
                'Subcategory': q['Subcategory'],
                'POS_Sale_Price': q['POS_Sale_Price'],
                'POS_Purchase_Price': q['POS_Purchase_Price'],
                'Stock_Qty': q['Stock_Qty'],
                'Verification_Evidence_URL': q['Verification_Evidence_URL'],
                'Reason': reason
            })

    manual_fields = [
        'Source_Row', 'POS_Product_Name', 'Brand', 'Category', 'Subcategory',
        'POS_Sale_Price', 'POS_Purchase_Price', 'Stock_Qty', 'Verification_Evidence_URL', 'Reason'
    ]
    with open(MANUAL_CSV, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=manual_fields)
        writer.writeheader()
        for r in final_manual_review_rows:
            writer.writerow(r)

    print(f"Saved {len(final_manual_review_rows)} unresolved queue items to {MANUAL_CSV}.")

    # 7. Complete Mathematical Inventory Audit
    wb_final = openpyxl.load_workbook(GOLD_XLSX, data_only=True)
    final_gold_by_sheet = {}
    for s in wb_final.sheetnames[:14]:
        final_gold_by_sheet[s] = count_actual_products(wb_final[s])
    total_final_gold = sum(final_gold_by_sheet.values())

    wb_hold = openpyxl.load_workbook(HOLD_XLSX, data_only=True)
    t1_hold = count_actual_products(wb_hold[wb_hold.sheetnames[0]])
    t2_hold = count_actual_products(wb_hold[wb_hold.sheetnames[1]])
    total_hold = t1_hold + t2_hold

    total_manual = len(final_manual_review_rows)
    grand_total = total_final_gold + total_hold + total_manual

    print("\n" + "=" * 80)
    print("FINAL 100% INVENTORY RECONCILIATION AUDIT")
    print("=" * 80)
    print(f"1. Verified Gold Master (HRZ_Gold_Master_Verified.xlsx):   {total_final_gold:4d} items")
    print(f"   -> Phase 1 Verified Baseline:                            122 items")
    print(f"   -> Phase 2 Newly Scraped & 100% Verified:               {len(verified_to_merge):4d} items")
    print(f"2. Client Hold Queue (Client_Hold_Generic_And_Unfound.xlsx): {total_hold:4d} items")
    print(f"   -> Tab 1 (N/A — Generic/Unbranded):                      {t1_hold:4d} items")
    print(f"   -> Tab 2 (Not Found Externally):                          {t2_hold:4d} items")
    print(f"3. Remaining Manual Review (scrape_manual_review.csv):       {total_manual:4d} items")
    print("-" * 80)
    print(f"GRAND TOTAL POS INVENTORY RECONCILED:                      {grand_total:4d} / 2502 items")
    print("=" * 80)

    print("\nCategory Distribution in Verified Gold Master:")
    for s_name, count in final_gold_by_sheet.items():
        print(f"  - {s_name:28s}: {count:4d} verified products")

    assert grand_total == 2502, f"Discrepancy detected: {grand_total} != 2502"
    print("\n>>> RECONCILIATION SUCCESS: 100.0% OF ALL 2,502 POS ITEMS ARE FULLY ACCOUNTED FOR! <<<")

if __name__ == '__main__':
    main()
