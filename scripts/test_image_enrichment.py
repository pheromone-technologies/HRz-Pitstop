import openpyxl
import json
import os

# 1. Load image mappings from Verified_Products_FINAL-1.xlsx
wb_old = openpyxl.load_workbook('Verified_Products_FINAL-1.xlsx', data_only=True)
sheet_old = wb_old['Verified Products']
headers_old = [str(c) for c in next(sheet_old.iter_rows(min_row=1, max_row=1, values_only=True))]

img_idx = headers_old.index('Product Image Filename / URL')
gal_idx = headers_old.index('Gallery Images')
name_idx = headers_old.index('Product Name')
sku_idx = headers_old.index('SKU')

old_img_map_by_name = {}
old_img_map_by_sku = {}

for row in sheet_old.iter_rows(min_row=2, values_only=True):
    pname = str(row[name_idx]).strip() if row[name_idx] else ''
    sku = str(row[sku_idx]).strip() if row[sku_idx] else ''
    img = row[img_idx] or ''
    gal = row[gal_idx] or ''
    if pname and (img or gal):
        old_img_map_by_name[pname.lower()] = (img, gal)
    if sku and sku != 'Missing SKU' and (img or gal):
        old_img_map_by_sku[sku] = (img, gal)

wb_old.close()

# 2. Also load images from data/products.json for variant-resolved images
with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

json_img_map = {}
for p in products:
    pname = p.get('name', '').strip().lower()
    img = p.get('image', '')
    gal = ', '.join(p.get('gallery', [])) if isinstance(p.get('gallery'), list) else (p.get('gallery') or '')
    if pname and (img or gal):
        json_img_map[pname] = (img, gal)
    if 'variants' in p:
        for v in p['variants']:
            vname = v.get('name', '').strip().lower()
            vsku = v.get('sku', '').strip()
            vimg = v.get('image', img)
            vgal = ', '.join(v.get('gallery', [])) if isinstance(v.get('gallery'), list) else gal
            if vname:
                json_img_map[vname] = (vimg, vgal)
            if vsku and vsku != 'Missing SKU':
                old_img_map_by_sku[vsku] = (vimg, vgal)

print(f"Old workbook mapped names: {len(old_img_map_by_name)}")
print(f"Old workbook mapped SKUs: {len(old_img_map_by_sku)}")
print(f"JSON mapped names: {len(json_img_map)}")

# 3. Test matching on Final_Master_Catalogue_VERIFIED.xlsx
wb_final = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', data_only=True)
sheet_final = wb_final['Master Catalogue']
headers_final = [str(c) for c in next(sheet_final.iter_rows(min_row=1, max_row=1, values_only=True))]

fn_name_idx = headers_final.index('Product Name')
fn_sku_idx = headers_final.index('SKU')

total_final = 0
found_img = 0
found_gal = 0

for row in sheet_final.iter_rows(min_row=2, values_only=True):
    total_final += 1
    pname = str(row[fn_name_idx]).strip() if row[fn_name_idx] else ''
    sku = str(row[fn_sku_idx]).strip() if row[fn_sku_idx] else ''
    
    img_gal = None
    if pname.lower() in old_img_map_by_name:
        img_gal = old_img_map_by_name[pname.lower()]
    elif sku in old_img_map_by_sku:
        img_gal = old_img_map_by_sku[sku]
    elif pname.lower() in json_img_map:
        img_gal = json_img_map[pname.lower()]
    
    if img_gal:
        if img_gal[0]: found_img += 1
        if img_gal[1]: found_gal += 1

wb_final.close()

print(f"Total Master Catalogue rows: {total_final}")
print(f"Rows matched with verified Image: {found_img}")
print(f"Rows matched with verified Gallery: {found_gal}")
