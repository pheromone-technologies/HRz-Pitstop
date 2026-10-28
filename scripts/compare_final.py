import openpyxl
import sys

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb_final = openpyxl.load_workbook(r'E:\coad\hrx website 2\HRZ_Pitstop_Client_Inventory_FINAL.xlsx', data_only=True)
wb_filled = openpyxl.load_workbook(r'E:\coad\hrx website 2\products data\HRZ_Pitstop_Client_Inventory_Filled.xlsx', data_only=True)
wb_gold = openpyxl.load_workbook(r'E:\coad\hrx website 2\HRZ_Gold_Master_Verified.xlsx', data_only=True)

print("=" * 70)
print("COMPARING HRZ_Pitstop_Client_Inventory_FINAL.xlsx")
print("=" * 70)

product_sheets = [s for s in wb_final.sheetnames if s not in ['LISTS', 'BIKE LISTS', 'MODEL DROPDOWNS']]

total_diff_rows = 0
for s in product_sheets:
    ws_f = wb_final[s]
    ws_old = wb_filled[s] if s in wb_filled.sheetnames else None
    
    rf = list(ws_f.iter_rows(values_only=True))
    ro = list(ws_old.iter_rows(values_only=True)) if ws_old else []
    
    diff_count = 0
    sample_diffs = []
    
    max_len = max(len(rf), len(ro))
    for i in range(max_len):
        row_f = rf[i] if i < len(rf) else None
        row_o = ro[i] if i < len(ro) else None
        if row_f != row_o:
            diff_count += 1
            if len(sample_diffs) < 3 and row_f and row_o:
                for c in range(min(len(row_f), len(row_o))):
                    if row_f[c] != row_o[c]:
                        sample_diffs.append((i, c, row_o[c], row_f[c]))
    
    total_diff_rows += diff_count
    print(f"Sheet: {s:<28} | Diff rows: {diff_count}")
    for idx, c, old_val, new_val in sample_diffs:
        print(f"   Row {idx} Col {c}: Old='{old_val}' -> New='{new_val}'")

print(f"\nTotal differing rows between FINAL and Filled: {total_diff_rows}")
