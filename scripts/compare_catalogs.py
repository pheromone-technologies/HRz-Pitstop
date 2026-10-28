import openpyxl
import json
from collections import Counter

wb_final = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet_final = wb_final['Master Catalogue']
headers_final = next(sheet_final.iter_rows(values_only=True))

final_by_name = {}
for row in sheet_final.iter_rows(values_only=True):
    d = dict(zip(headers_final, row))
    pname = d.get('Product Name')
    if pname and pname not in final_by_name:
        final_by_name[pname] = d
wb_final.close()

wb_old = openpyxl.load_workbook('Verified_Products_FINAL-1.xlsx', read_only=True)
sheet_old = wb_old['Verified Products']
headers_old = next(sheet_old.iter_rows(values_only=True))

old_by_name = {}
for row in sheet_old.iter_rows(values_only=True):
    d = dict(zip(headers_old, row))
    pname = d.get('Product Name')
    if pname and pname not in old_by_name:
        old_by_name[pname] = d
wb_old.close()

# 1. Compare Category / Subcategory changes for items present in both
cat_changes = []
subcat_changes = []
for pname, old_d in old_by_name.items():
    if pname in final_by_name:
        fn_d = final_by_name[pname]
        if old_d.get('Category') != fn_d.get('Category'):
            cat_changes.append((pname, old_d.get('Category'), fn_d.get('Category')))
        if old_d.get('Subcategory') != fn_d.get('Subcategory'):
            subcat_changes.append((pname, old_d.get('Subcategory'), fn_d.get('Subcategory')))

print(f'Items with Category changed: {len(cat_changes)}')
for c in cat_changes[:15]:
    print(f'  {c[0][:40]}: "{c[1]}" -> "{c[2]}"')

print(f'\nItems with Subcategory changed: {len(subcat_changes)}')
for s in subcat_changes[:15]:
    print(f'  {s[0][:40]}: "{s[1]}" -> "{s[2]}"')

# 2. Inspect the new items in Final_Master
new_in_final = [d for pname, d in final_by_name.items() if pname not in old_by_name]
print(f'\nNew items in Final_Master: {len(new_in_final)}')
new_status = Counter(d.get('Verification') for d in new_in_final)
print('New items Verification status:')
for k, v in new_status.items():
    print(f'  {k}: {v}')

new_confidence = Counter(d.get('Confidence') for d in new_in_final)
print('\nNew items Confidence:')
for k, v in new_confidence.items():
    print(f'  {k}: {v}')

# Sample of new items
print('\nSample 10 new items:')
for d in new_in_final[:10]:
    print(f"  Name: {d.get('Product Name')} | Cat: {d.get('Category')} | Subcat: {d.get('Subcategory')} | Conf: {d.get('Confidence')} | Price: {d.get('Sale Price')}")
