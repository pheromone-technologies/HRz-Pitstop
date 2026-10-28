import json
import openpyxl

with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

# Also check Master Catalogue what Size / Variant values exist
wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))
name_idx = headers.index('Product Name')
size_idx = headers.index('Size / Variant')
cat_idx = headers.index('Category')
subcat_idx = headers.index('Subcategory')

master_sizes = {}
for row in sheet.iter_rows(min_row=2, values_only=True):
    pname = str(row[name_idx]).strip() if row[name_idx] else ''
    sz = str(row[size_idx]).strip() if row[size_idx] else ''
    if pname:
        master_sizes[pname.lower()] = sz
wb.close()

sized = [p for p in products if p.get('availableSizes') or p.get('sizes') or p.get('allSizes')]
print(f"Total products in products.json: {len(products)}")
print(f"Products with sizes in products.json: {len(sized)}")

# Check non-wearable or non-sized products
false_sizes = []
for p in sized:
    name_l = p['name'].lower()
    cat_l = (p.get('category') or '').lower()
    subcat_l = (p.get('subcategory') or '').lower()
    
    # Check what Master Catalogue has
    orig_size = master_sizes.get(name_l, '')
    
    # Electronic devices, bluetooth, intercom, visor, pinlock, mounts, etc.
    is_non_wearable = any(k in name_l for k in [
        'bluetooth', 'intercom', 'visor', 'spoiler', 'pinlock', 'mount',
        'charger', 'lock', 'cable', 'cover', 'strap', 'battery', 'mic',
        'speaker', 'clamp', 'screw', 'padding', 'mechanism', 'plate'
    ])
    
    # Or if category is Helmet Accessories / Luggage / Cleaning
    if is_non_wearable or 'accessories' in subcat_l or 'clean' in cat_l or 'luggage' in cat_l:
        false_sizes.append((p, orig_size))

print(f"\nNon-sized products falsely assigned sizes: {len(false_sizes)}")
for p, orig_sz in false_sizes:
    print(f"  Name: {p['name']} | Cat: {p.get('category')} / {p.get('subcategory')} | Master Size: '{orig_sz}' | Injected sizes: {p.get('sizes')}")
