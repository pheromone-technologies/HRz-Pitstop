import json
import os
import io
import requests
import re
from PIL import Image
import numpy as np
import openpyxl

# Import our standardizer
from image_standardizer import standardize_and_save

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
}

def download_and_standardize(url, out_path):
    print(f"Downloading from {url} -> {out_path} ...")
    try:
        r = requests.get(url, headers=HEADERS, timeout=20)
        if r.status_code != 200:
            print(f"FAILED status {r.status_code} for {url}")
            return False
        im = Image.open(io.BytesIO(r.content))
        success = standardize_and_save(im, out_path)
        if success:
            print(f"  -> Successfully standardized and saved to {out_path}")
            return True
        else:
            print(f"  -> Standardization returned False for {out_path}")
            return False
    except Exception as e:
        print(f"ERROR downloading {url}: {e}")
        return False

# 1. Target Downloads & Standardizations
tasks = [
    # User Target 1: SBH-25 Armada Blue Full Face
    {
        'name': 'SBH-25 Armada Blue Full Face',
        'source': 'scratch/sbh25_armada_blue_real.webp' if os.path.exists('scratch/sbh25_armada_blue_real.webp') else 'https://cdn.shopify.com/s/files/1/0694/9673/3873/files/steelbird_sbh_25_sharp_armada_blue_full_face_helmet_1.webp?v=1781156327',
        'dest': 'images_processed/catalog/steelbird/sbh_25_sharp_armada_blue.webp'
    },
    # User Target 2: Axor Apex Venomous distinct colorways
    {
        'name': 'Axor Apex Venomous Black Grey',
        'source': 'https://cdn.shopify.com/s/files/1/0694/9673/3873/files/apex_venomous_d_v_black_grey__1.webp?v=1781154176',
        'dest': 'images_processed/catalog/axor/axor_apex_venomous_black_grey_m.webp',
        'alias': 'images_processed/catalog/axor/axor_apex_venomous_black_grey_m_1.webp'
    },
    {
        'name': 'Axor Apex Venomous Black Red',
        'source': 'https://cdn.shopify.com/s/files/1/0694/9673/3873/files/axor_apex_venomous_black_red_full_face_helmet_1_1792652712805_331.webp?v=1792652731',
        'dest': 'images_processed/catalog/axor/axor_apex_venomous_black_red_l.webp',
        'alias': 'images_processed/catalog/axor/axor_apex_venomous_black_red_l_1.webp'
    },
    {
        'name': 'Axor Apex Venomous Black Neon Blue',
        'source': 'https://cdn.shopify.com/s/files/1/0694/9673/3873/files/apex_venomous_d_v_black_neon_blue__1.webp?v=1781154179',
        'dest': 'images_processed/catalog/axor/axor_apex_venomous_black_blue_xl.webp',
        'alias': 'images_processed/catalog/axor/axor_apex_venomous_black_blue_xl_1.webp'
    },
    # User Target 3: SBH-61 SXE Fluro-Red Full Face
    {
        'name': 'Steelbird SBH-61 SXE SP3 Fluro-Red',
        'source': 'https://cdn.shopify.com/s/files/1/0694/9673/3873/files/steelbird_sbh_61_sxe_sp3_black_red_full_face_helmet_1.webp?v=1781240075',
        'dest': 'images_processed/catalog/steelbird/sbh-61_d_v_sxe_tpe_st2_glsy_blk_flurored_helmet_600mm.webp'
    },
    # Comprehensive Fix 4: TT Revo Plain Black
    {
        'name': 'KYT TT Revo Plain Black',
        'source': 'https://cdn.shopify.com/s/files/1/1459/5894/files/Y6TRX0W6-KYTTT-REVOPLAINBLACK_8.png?v=1743160625',
        'dest': 'images_processed/catalog/kyt/tt_revo_plain_black_l.webp'
    },
    # Comprehensive Fix 5: TT Revo Electron Matt Grey Red
    {
        'name': 'KYT TT Revo Electron Matt Grey Red',
        'source': 'https://superbikestore.in/cdn/shop/files/1_31a6d904-97e0-43ff-9b80-15ff9639cbb3.png?v=1790769944',
        'dest': 'images_processed/catalog/kyt/tt_revo_electron_matt_grey_red_l.webp'
    },
    # Comprehensive Fix 6: SBA-20 / SB-106 Strange Jazz Blue Flip-Up
    {
        'name': 'Steelbird SBA-20 Strange Jazz Blue',
        'source': 'https://rukminim2.flixcart.com/image/832/832/xif0q/helmet/w/q/n/sba-20-strange-isi-certified-flip-up-graphic-helmet-for-men-and-original-imaheb6w6tpeyu5t.jpeg',
        'dest': 'images_processed/catalog/steelbird/sba_20_strange_jazz_blue.webp'
    },
    # Comprehensive Fix 7: SBH-57 Plain Black
    {
        'name': 'Steelbird SBH-57 Fighter Plain Black',
        'source': 'https://cdn.shopify.com/s/files/1/0938/2487/6828/files/6676c58f624860T4A3902.png?v=1759939955',
        'dest': 'images_processed/catalog/steelbird/sbh-57_d_v_fighter_glossy_black_helmet.webp'
    }
]

print("\n=== STEP 1: Processing All Images ===")
for t in tasks:
    src = t['source']
    dest = t['dest']
    if os.path.exists(src):
        print(f"Standardizing local file {src} -> {dest} ...")
        im = Image.open(src)
        standardize_and_save(im, dest)
    else:
        download_and_standardize(src, dest)
    if 'alias' in t:
        alias = t['alias']
        im = Image.open(dest)
        standardize_and_save(im, alias)

print("\n=== STEP 2: Updating Catalog JSONs ===")
catalog_files = ['data/products.json', 'test design site/data/products.json']

for cat_file in catalog_files:
    with open(cat_file, 'r', encoding='utf-8') as f:
        prods = json.load(f)

    # 1. SBH-25 SHARP (D.V) GLOSSY ARMADA BLUE HELMET
    for p in prods:
        if p['id'] == 'hrz-sbh-25-sharp-d-v-glossy-armada-blue-helmet-2618':
            p['image'] = 'images_processed/catalog/steelbird/sbh_25_sharp_armada_blue.webp'
            p['gallery'] = ['images_processed/catalog/steelbird/sbh_25_sharp_armada_blue.webp']
        # Also ensure SBH 25 SV SHARP Glossy Black uses plain black image, NOT Albus Red
        if p['id'] == 'hrz-sbh-25-sv-sharp-helmet-2602':
            p['image'] = 'images_processed/catalog/steelbird/sbh_25_sharp_glossy_black.webp'
            p['gallery'] = ['images_processed/catalog/steelbird/sbh_25_sharp_glossy_black.webp']

        # Update variants on SBH-25 products
        if 'sbh-25' in p['id'] and p.get('variants'):
            for v in p['variants']:
                c_up = (v.get('color') or '').upper()
                if v.get('id') == 'hrz-sbh-25-sharp-d-v-glossy-armada-blue-helmet-2618' or 'ARMADA' in c_up:
                    v['image'] = 'images_processed/catalog/steelbird/sbh_25_sharp_armada_blue.webp'
                    v['gallery'] = ['images_processed/catalog/steelbird/sbh_25_sharp_armada_blue.webp']
                elif v.get('id') == 'hrz-sbh-25-d-v-sharp-albus-glossy-black-red-helmet-2591' or 'BLACK/RED' in c_up:
                    v['image'] = 'images_processed/catalog/steelbird/sbh-25_d_v_sharp_albus_glossy_black_red_helmet.webp'
                    v['gallery'] = ['images_processed/catalog/steelbird/sbh-25_d_v_sharp_albus_glossy_black_red_helmet.webp']
                elif v.get('id') == 'hrz-sbh-25-d-v-sharp-albus-glossy-matt-blk-helmet-2593' or 'MATT' in c_up:
                    v['image'] = 'images_processed/catalog/steelbird/sbh_25_sharp_albus_matt_black.webp'
                    v['gallery'] = ['images_processed/catalog/steelbird/sbh_25_sharp_albus_matt_black.webp']
                elif v.get('id') == 'hrz-sbh-25-sharp-d-v-glossy-black-helmet-2604' or c_up == 'GLOSSY BLACK':
                    v['image'] = 'images_processed/catalog/steelbird/sbh_25_sharp_glossy_black.webp'
                    v['gallery'] = ['images_processed/catalog/steelbird/sbh_25_sharp_glossy_black.webp']

        # 2. Axor Apex Venomous
        if p['id'] == 'hrz-axor-apex-venomous-black-grey-m-0166':
            p['image'] = 'images_processed/catalog/axor/axor_apex_venomous_black_grey_m.webp'
            p['gallery'] = ['images_processed/catalog/axor/axor_apex_venomous_black_grey_m.webp']
        if p['id'] == 'hrz-axor-apex-venomous-black-red-l-0167':
            p['image'] = 'images_processed/catalog/axor/axor_apex_venomous_black_red_l.webp'
            p['gallery'] = ['images_processed/catalog/axor/axor_apex_venomous_black_red_l.webp']
        if p['id'] == 'hrz-axor-apex-venomous-black-blue-xl-0168':
            p['image'] = 'images_processed/catalog/axor/axor_apex_venomous_black_blue_xl.webp'
            p['gallery'] = ['images_processed/catalog/axor/axor_apex_venomous_black_blue_xl.webp']

        # Update Axor Apex Venomous variants if present
        if 'venomous' in p['id'] and p.get('variants'):
            for v in p['variants']:
                c_up = (v.get('color') or '').upper()
                if v.get('id') == 'hrz-axor-apex-venomous-black-grey-m-0166' or 'GREY' in c_up:
                    v['image'] = 'images_processed/catalog/axor/axor_apex_venomous_black_grey_m.webp'
                    v['gallery'] = ['images_processed/catalog/axor/axor_apex_venomous_black_grey_m.webp']
                elif v.get('id') == 'hrz-axor-apex-venomous-black-red-l-0167' or 'RED' in c_up:
                    v['image'] = 'images_processed/catalog/axor/axor_apex_venomous_black_red_l.webp'
                    v['gallery'] = ['images_processed/catalog/axor/axor_apex_venomous_black_red_l.webp']
                elif v.get('id') == 'hrz-axor-apex-venomous-black-blue-xl-0168' or 'BLUE' in c_up:
                    v['image'] = 'images_processed/catalog/axor/axor_apex_venomous_black_blue_xl.webp'
                    v['gallery'] = ['images_processed/catalog/axor/axor_apex_venomous_black_blue_xl.webp']

        # 3. Steelbird SBH-61
        if p['id'] == 'hrz-sbh-61-d-v-sxe-tpe-st2-glsy-blk-flurored-helmet-60-2640':
            p['image'] = 'images_processed/catalog/steelbird/sbh-61_d_v_sxe_tpe_st2_glsy_blk_flurored_helmet_600mm.webp'
            p['gallery'] = ['images_processed/catalog/steelbird/sbh-61_d_v_sxe_tpe_st2_glsy_blk_flurored_helmet_600mm.webp']
        if p['id'] in ['hrz-sbh-61-d-v-sxe-afs-edition-glsy-blk-helmet-620mm-s-2630', 'hrz-sbh-61-d-v-sxe-afs-glossy-black-helmet-560mm-2638']:
            p['image'] = 'images_processed/catalog/steelbird/sbh_61_sxe_afs_gloss_black.webp'
            p['gallery'] = ['images_processed/catalog/steelbird/sbh_61_sxe_afs_gloss_black.webp']

        # 4. KYT TT Revo Plain Black
        if p['id'] == 'hrz-tt-revo-plain-black-l-3049':
            p['image'] = 'images_processed/catalog/kyt/tt_revo_plain_black_l.webp'
            p['gallery'] = ['images_processed/catalog/kyt/tt_revo_plain_black_l.webp']

        # 5. KYT TT Revo Electron Matt Grey Red
        if p['id'] == 'hrz-tt-revo-electron-matt-grey-red-l-3023':
            p['image'] = 'images_processed/catalog/kyt/tt_revo_electron_matt_grey_red_l.webp'
            p['gallery'] = ['images_processed/catalog/kyt/tt_revo_electron_matt_grey_red_l.webp']

        # 6. SBA-20 / SB-106 Strange Jazz Blue
        if p['id'] == 'hrz-sba-20-sb106-d-v-strange-glsy-blk-jazz-blue-helmet-2561':
            p['image'] = 'images_processed/catalog/steelbird/sba_20_strange_jazz_blue.webp'
            p['gallery'] = ['images_processed/catalog/steelbird/sba_20_strange_jazz_blue.webp']

        # 7. SMK Agnar Cyborg GL263
        if p['id'] == 'hrz-smk-agnar-cyborgglo-bla-gre-red-gl263m-2677':
            p['image'] = 'images_processed/catalog/smk/smk_agnar_cyborg_graphic_gl263_lpf_1.webp'
            p['gallery'] = ['images_processed/catalog/smk/smk_agnar_cyborg_graphic_gl263_lpf_1.webp']

        # 8. Vega Helmet Axor
        if p['id'] == 'hrz-vega-helmet-axor-3110':
            p['image'] = 'images_processed/catalog/vega/vega_bolt_gloss_black.webp'
            p['gallery'] = ['images_processed/catalog/vega/vega_bolt_gloss_black.webp']

        # 9. SBH-57 Fighter Glossy Black
        if p['id'] == 'hrz-sbh-57-d-v-fighter-glossy-black-helmet-2624':
            p['image'] = 'images_processed/catalog/steelbird/sbh-57_d_v_fighter_glossy_black_helmet.webp'
            p['gallery'] = ['images_processed/catalog/steelbird/sbh-57_d_v_fighter_glossy_black_helmet.webp']

    with open(cat_file, 'w', encoding='utf-8') as f:
        json.dump(prods, f, indent=2, ensure_ascii=False)
    print(f"Saved updated catalog to {cat_file}")

print("\n=== STEP 3: Updating Master Excel Sheet ===")
wb_path = 'Final_Master_Catalogue_VERIFIED.xlsx'
if os.path.exists(wb_path):
    wb = openpyxl.load_workbook(wb_path)
    ws = wb.active
    # Map header columns
    headers = [cell.value for cell in ws[1]]
    id_col = headers.index('ID') + 1 if 'ID' in headers else 1
    img_col = headers.index('Image URL') + 1 if 'Image URL' in headers else (headers.index('IMAGE') + 1 if 'IMAGE' in headers else None)
    
    # Load products to get latest images
    with open('data/products.json', 'r', encoding='utf-8') as f:
        prods = json.load(f)
    id_map = {p['id']: p['image'] for p in prods}

    if img_col:
        for row in range(2, ws.max_row + 1):
            pid = ws.cell(row=row, column=id_col).value
            if pid in id_map:
                ws.cell(row=row, column=img_col).value = id_map[pid]
        wb.save(wb_path)
        print("Updated and saved Final_Master_Catalogue_VERIFIED.xlsx")

print("\n=== ALL COMPREHENSIVE HELMET FIXES APPLIED SUCCESSFULLY ===")
