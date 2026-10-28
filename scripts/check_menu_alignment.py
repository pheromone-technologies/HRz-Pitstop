import json
import re
from collections import Counter

# 1. Menu links in js/menu.js
with open('js/menu.js', 'r', encoding='utf-8') as f:
    menu_content = f.read()

# Find all category= and subcategory= in menu.js
menu_cats = set(re.findall(r'category=([^&"\']+)', menu_content))
menu_subcats = set(re.findall(r'subcategory=([^&"\']+)', menu_content))

import urllib.parse
menu_cats = {urllib.parse.unquote(c) for c in menu_cats}
menu_subcats = {urllib.parse.unquote(s) for s in menu_subcats}

print("Menu Categories from js/menu.js:")
for c in sorted(menu_cats):
    print(f"  {c}")

# 2. Categories in products.json
with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

prod_cats = Counter(p.get('category', '') for p in products)
print("\nCurrent Categories in data/products.json:")
for c, cnt in prod_cats.most_common():
    in_menu = "YES" if c in menu_cats else "NO"
    print(f"  {c}: {cnt} (In Menu: {in_menu})")

# 3. Check categories in Final_Master_Catalogue_VERIFIED.xlsx
import openpyxl
wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))
final_cats = Counter()
for row in sheet.iter_rows(values_only=True):
    d = dict(zip(headers, row))
    final_cats[d.get('Category', '')] += 1
wb.close()

print("\nCategories in Final_Master_Catalogue_VERIFIED.xlsx:")
for c, cnt in final_cats.most_common():
    in_menu = "YES" if c in menu_cats else "NO"
    print(f"  {c}: {cnt} (In Menu: {in_menu})")
