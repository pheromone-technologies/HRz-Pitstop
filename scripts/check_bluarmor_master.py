import openpyxl

wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))

for row in sheet.iter_rows(values_only=True):
    d = dict(zip(headers, row))
    pname = str(d.get('Product Name') or '').lower()
    if 'bluarmor' in pname or 'hs1' in pname:
        print(f"Name: {d.get('Product Name')} | Size/Variant: '{d.get('Size / Variant')}' | Cat: {d.get('Category')} | Subcat: {d.get('Subcategory')}")
wb.close()
