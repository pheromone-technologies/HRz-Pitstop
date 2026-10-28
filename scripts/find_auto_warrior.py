import os
import openpyxl

print("Searching for 'Auto Warrior' across directories...")
for root, dirs, files in os.walk('.'):
    if any(k in root for k in ['.git', '.venv']):
        continue
    for f in files:
        if 'warrior' in f.lower() or 'autowarrior' in f.lower():
            print(f"  File: {os.path.join(root, f)}")

# Search inside Excel files for Auto Warrior
excel_files = [f for f in os.listdir('.') if f.endswith('.xlsx')]
for ef in excel_files:
    try:
        wb = openpyxl.load_workbook(ef, read_only=True)
        found = 0
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            for row in sheet.iter_rows(values_only=True):
                row_str = ' '.join(str(c) for c in row if c is not None).lower()
                if 'auto warrior' in row_str or 'autowarrior' in row_str:
                    found += 1
        if found:
            print(f"  Found {found} occurrences in {ef}")
        wb.close()
    except Exception as e:
        pass
