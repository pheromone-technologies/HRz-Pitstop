import openpyxl
import json

wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))

for row in sheet.iter_rows(values_only=True):
    row_str = ' '.join(str(c) for c in row if c is not None)
    if 'auto warrior' in row_str.lower():
        d = dict(zip(headers, row))
        print("Master Row:")
        print(f"  SKU: {d.get('SKU')}")
        print(f"  Name: {d.get('Product Name')}")
        print(f"  Brand: {d.get('Brand')}")
        print(f"  Category: {d.get('Category')} / {d.get('Subcategory')}")
        print(f"  Evidence: {d.get('Evidence Link')}")
        print(f"  Img: {d.get('Product Image Filename / URL')}")
wb.close()

with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

print("\nMatching in products.json:")
for p in products:
    p_str = json.dumps(p).lower()
    if 'auto warrior' in p_str:
        print(f"  ID: {p['id']} | Name: {p['name']} | Img: {p.get('image')}")
