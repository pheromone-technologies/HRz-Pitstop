import openpyxl
import json
import os

wb = openpyxl.load_workbook('Final_Master_Catalogue_VERIFIED.xlsx', read_only=True)
sheet = wb['Master Catalogue']
headers = next(sheet.iter_rows(values_only=True))

final_rows = []
for row in sheet.iter_rows(values_only=True):
    final_rows.append(dict(zip(headers, row)))
wb.close()

with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

prod_by_name = {p['name'].strip().lower(): p for p in products}
prod_by_sku = {p.get('sku', '').strip(): p for p in products if p.get('sku')}

print(f"Total rows in Final_Master: {len(final_rows)}")
print(f"Total current products in data/products.json: {len(products)}")

matched_by_name = 0
matched_by_sku = 0
unmatched = []

for r in final_rows:
    name = (r.get('Product Name') or '').strip().lower()
    sku = (r.get('SKU') or '').strip()
    if name in prod_by_name:
        matched_by_name += 1
    elif sku and sku != 'Missing SKU' and sku in prod_by_sku:
        matched_by_sku += 1
    else:
        unmatched.append(r)

print(f"Matched by Name: {matched_by_name}")
print(f"Matched by SKU: {matched_by_sku}")
print(f"Unmatched (new/not currently in products.json): {len(unmatched)}")

# Breakdown of unmatched items
from collections import Counter
unmatched_status = Counter(r.get('Verification') for r in unmatched)
print("\nUnmatched items by Verification Status:")
for k, v in unmatched_status.items():
    print(f"  {k}: {v}")

unmatched_conf = Counter(r.get('Confidence') for r in unmatched)
print("\nUnmatched items by Confidence:")
for k, v in unmatched_conf.items():
    print(f"  {k}: {v}")
