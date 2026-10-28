#!/usr/bin/env python3
"""
HRz Pitstop: Category & Subcategory Normalization Script
Cleans up misplaced items and enriches subcategories across data/products.json and Verified_Products_FINAL-1.xlsx
"""

import os
import sys
import json
import re
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"E:\coad\hrx website 2"
EXCEL_FILE = os.path.join(BASE_DIR, "Verified_Products_FINAL-1.xlsx")
PRODUCTS_JSON = os.path.join(BASE_DIR, "data", "products.json")

def detect_category_and_subcategory(name, current_cat, current_subcat):
    n = name.lower()
    
    # 1. Bike Windshields/Visors mistakenly placed in Helmets
    if any(k in n for k in ['xpulse touring visor', 'r15 v4 visor', 'r15 v4 bolt', 'dominar visor', 'dominar touring', 'mt metal visor', 'polycarbon black', 'touring 1.0 visor']):
        return 'Body Parts & Panels', 'Windshields & Visors'
        
    # 2. Headlight Guards mistakenly placed in Helmets
    if 'headlight guard' in n or 'headlight grill' in n or 'head light grill' in n:
        return 'Crash Protection', 'Headlight Grills'
        
    # 3. Sleeves & Balaclavas
    if 'sleeves' in n or 'balaclava' in n or 'balclavaa' in n:
        return 'Riding Jackets & Apparel', 'Headwear & Balaclavas'
        
    # 4. Goggles & Eyewear
    if 'goggles' in n or 'riding glasses' in n or 'mask black tint' in n:
        return 'Riding Jackets & Apparel', 'Riding Eyewear'
        
    # 5. Helmets
    if current_cat == 'Helmets' or 'helmet' in n:
        # Determine helmet subcategory
        if any(k in n for k in ['visor', 'spoiler', 'pinlock', 'insert', 'ratchet', 'pivot kit', 'short vis']):
            return 'Helmets', 'Visors & Spoilers'
        if any(k in n for k in ['bluetooth', 'intercom', 'chinmount', 'fitting', 'anti fog', 'antifog']):
            return 'Helmets', 'Helmet Accessories'
        if any(k in n for k in ['open face', 'openface', 'o/f', 'jet', 'buds open', 'cruiserwith peak', 'half face', 'half head']):
            return 'Helmets', 'Open Face'
        if any(k in n for k in ['modular', 'gullwing', 'boolean', 'evo dv', 'flip up']):
            return 'Helmets', 'Modular'
        if any(k in n for k in ['dirt', 'allterra', 'v cross', 'x-cross', 'x cross', 'motocross', 'off road', 'offroad']):
            return 'Helmets', 'Off-Road'
        return 'Helmets', 'Full Face'
        
    # 6. Riding Jackets & Apparel
    if current_cat in ['Riding Jackets & Apparel', 'Riding Pants', 'Gloves', 'Boots & Footwear']:
        if 'jacket' in n or 'jersey' in n:
            return 'Riding Jackets & Apparel', 'Jackets'
        if 'glove' in n:
            return 'Gloves', 'Riding Gloves'
        if 'pant' in n or 'trouser' in n:
            return 'Riding Pants', 'Riding Pants'
        if 'boot' in n or 'shoe' in n:
            return 'Boots & Footwear', 'Boots & Footwear'
        return current_cat, current_subcat or 'Apparel'
        
    # 7. Crash Protection
    if current_cat == 'Crash Protection' or any(k in n for k in ['crash guard', 'bash plate', 'slider', 'radiator guard']):
        if 'bash plate' in n or 'skid plate' in n:
            return 'Crash Protection', 'Bash Plates'
        if 'radiator' in n:
            return 'Crash Protection', 'Radiator Guards'
        if 'slider' in n or 'bobbin' in n:
            return 'Crash Protection', 'Frame Sliders'
        if 'headlight' in n:
            return 'Crash Protection', 'Headlight Grills'
        return 'Crash Protection', 'Crash Guards'
        
    # 8. Luggage & Storage
    if current_cat == 'Luggage & Storage' or any(k in n for k in ['saddlebag', 'saddle bag', 'tank bag', 'tail bag', 'backpack', 'top box', 'pannier']):
        if 'saddle bag' in n or 'saddlebag' in n:
            return 'Luggage & Storage', 'Saddlebags'
        if 'tank bag' in n:
            return 'Luggage & Storage', 'Tank Bags'
        if 'tail bag' in n:
            return 'Luggage & Storage', 'Tail Bags'
        if 'backpack' in n:
            return 'Luggage & Storage', 'Backpacks'
        if 'top box' in n or 'top case' in n or 'box' in n:
            return 'Luggage & Storage', 'Top Boxes'
        if 'stay' in n or 'rack' in n or 'carrier' in n:
            return 'Luggage & Storage', 'Luggage Racks & Stays'
        return 'Luggage & Storage', 'Touring Bags'
        
    # 9. Lighting & Electrical
    if current_cat == 'Lighting & Electrical' or any(k in n for k in ['fog light', 'fog lamp', 'indicator', 'headlight', 'hazard', 'auxiliary']):
        if 'fog' in n or 'auxiliary' in n:
            return 'Lighting & Electrical', 'Fog & Aux Lights'
        if 'indicator' in n or 'flasher' in n or 'hazard' in n:
            return 'Lighting & Electrical', 'Indicators & Hazard'
        if 'headlight' in n:
            return 'Lighting & Electrical', 'Headlights'
        if 'tail' in n:
            return 'Lighting & Electrical', 'Tail Lights'
        return 'Lighting & Electrical', 'Lighting Accessories'
        
    return current_cat, current_subcat

def main():
    print("=" * 80)
    print("NORMALIZING CATEGORIES & SUBCATEGORIES ACROSS CATALOG")
    print("=" * 80)

    # 1. Update data/products.json
    with open(PRODUCTS_JSON, 'r', encoding='utf-8') as f:
        products = json.load(f)

    updated_count = 0
    for p in products:
        old_cat = p.get('category', '')
        old_sub = p.get('subcategory', '')
        new_cat, new_sub = detect_category_and_subcategory(p['name'], old_cat, old_sub)
        
        if new_cat != old_cat or new_sub != old_sub:
            p['category'] = new_cat
            p['subcategory'] = new_sub
            updated_count += 1

    with open(PRODUCTS_JSON, 'w', encoding='utf-8') as f:
        json.dump(products, f, indent=2)
    print(f"Updated {updated_count} products in data/products.json with normalized categories & subcategories.")

    # 2. Update Verified_Products_FINAL-1.xlsx
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    name_col = headers.index('Product Name') + 1
    cat_col = headers.index('Category') + 1
    sub_col = headers.index('Subcategory') + 1

    excel_updated = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, name_col).value or ''
        c_cat = ws.cell(r, cat_col).value or ''
        c_sub = ws.cell(r, sub_col).value or ''
        
        n_cat, n_sub = detect_category_and_subcategory(name, c_cat, c_sub)
        if n_cat != c_cat or n_sub != c_sub:
            ws.cell(r, cat_col, n_cat)
            ws.cell(r, sub_col, n_sub)
            excel_updated += 1

    wb.save(EXCEL_FILE)
    print(f"Updated {excel_updated} rows in {EXCEL_FILE} with normalized categories & subcategories.")

    # Show new category distribution
    from collections import Counter
    cats = Counter(p['category'] for p in products)
    print("\nNew Storefront Category Distribution:")
    for k, v in cats.most_common():
        print(f"  {k}: {v} products")

    helmets = [p for p in products if p['category'] == 'Helmets']
    helmet_subs = Counter(p['subcategory'] for p in helmets)
    print(f"\nHelmets Subcategory Distribution (Total {len(helmets)}):")
    for k, v in helmet_subs.items():
        print(f"  {k}: {v}")

    print("=" * 80)

if __name__ == '__main__':
    main()
